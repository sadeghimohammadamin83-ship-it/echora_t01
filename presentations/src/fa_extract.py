# Extract every distinct text run from the two English decks (for translation to Persian).
import json, re, sys
from pptx import Presentation
def runs(prs):
    for slide in prs.slides:
        for sh in slide.shapes:
            frames = []
            if sh.has_text_frame: frames.append(sh.text_frame)
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    for c in row.cells: frames.append(c.text_frame)
            for tf in frames:
                for p in tf.paragraphs:
                    for r in p.runs: yield r
    for slide in prs.slides:
        for sh in slide.slide_layout.placeholders: pass
keep = {}
for f in sys.argv[1:]:
    for r in runs(Presentation(f)):
        t = r.text
        if not t.strip(): continue
        if re.fullmatch(r"[\d\s.,×x–\-+%/()≥≤≈~mcmM²:]+", t): continue
        if t.startswith("http"): continue
        keep[t] = keep.get(t, 0) + 1
json.dump(sorted(keep), open("src/fa_strings.json", "w"), ensure_ascii=False, indent=0)
print(len(keep), sum(len(k) for k in keep))
