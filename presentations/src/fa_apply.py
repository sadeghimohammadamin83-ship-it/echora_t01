"""Build Persian (RTL) versions of the two decks from the English .pptx files.

Usage: python3 fa_apply.py in.pptx out.pptx
Replaces every text run listed in fa_map.FA (keyed by fa_strings.json), sets
paragraphs that contain Persian to right-to-left, right-aligns left-aligned
text, mirrors tables, and switches all text to the Vazirmatn font.
"""
import json, re, sys, os
from pptx import Presentation
from pptx.oxml.ns import qn
from lxml import etree

HERE = os.path.dirname(__file__)
sys.path.insert(0, HERE)
from fa_map import FA, EXTRA, FOOTER

STRINGS = json.load(open(os.path.join(HERE, "fa_strings.json")))
TR = {STRINGS[i]: t for i, t in FA.items()}
TR.update(EXTRA)
TR.update(FOOTER)
FONT = "Vazirmatn"
SCALE = 0.92          # Vazirmatn sets wider/taller than Arial
PERSIAN = re.compile(r"[؀-ۿ]")


def set_font(rPr):
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", FONT)
    # keep schema order: latin, ea, cs must follow fill elements — move them to the end
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        rPr.remove(el)
        rPr.append(el)
    for tag in ("a:hlinkClick",):
        el = rPr.find(qn(tag))
        if el is not None:
            rPr.remove(el)
            rPr.append(el)
    rPr.set("lang", "fa-IR")
    if rPr.get("sz"):
        rPr.set("sz", str(int(int(rPr.get("sz")) * SCALE)))


def do_frame(tf, width_in=99.0):
    wide = width_in >= 2.2   # narrow boxes are labels beside symbols/dimension lines: keep their anchor side
    frame_fa = any(PERSIAN.search(TR.get(r.text, r.text)) for p in tf.paragraphs for r in p.runs)
    for p in tf.paragraphs:
        persian = False
        for r in p.runs:
            if r.text in TR:
                r.text = TR[r.text]
            if "↔" in r.text:
                r.text = r.text.replace("↔", "⇄")   # ↔ falls back to a colour-emoji glyph in Persian runs
            if PERSIAN.search(r.text):
                persian = True
            rPr = r._r.get_or_add_rPr()
            set_font(rPr)
            if "spc" in rPr.attrib:      # letter-spacing breaks Persian joining
                del rPr.attrib["spc"]
        pPr = p._p.get_or_add_pPr()
        if persian:
            pPr.set("rtl", "1")
        big = any((r.font.size or 0) >= 177800 for r in p.runs)   # >= 14 pt (stat values)
        if wide and (persian or big or frame_fa) and pPr.get("algn") in (None, "l") and any(r.text.strip() for r in p.runs):
            pPr.set("algn", "r")


def do_shapes(shapes):
    for sh in shapes:
        if sh.shape_type == 6:  # group
            do_shapes(sh.shapes)
        if sh.has_text_frame:
            do_frame(sh.text_frame, (sh.width or 0) / 914400 if sh.width else 99.0)
        if getattr(sh, "has_table", False) and sh.has_table:
            tblPr = sh._element.find(".//" + qn("a:tblPr"))
            if tblPr is not None:
                tblPr.set("rtl", "1")
            for row in sh.table.rows:
                for c in row.cells:
                    do_frame(c.text_frame)


def fix_content_types(path):
    """python-pptx can drop the 'jpg' default content type; add it back if any .jpg part exists."""
    import zipfile, shutil
    zin = zipfile.ZipFile(path)
    ct = zin.read("[Content_Types].xml").decode()
    if any(n.endswith(".jpg") for n in zin.namelist()) and 'Extension="jpg"' not in ct:
        ct = ct.replace("<Default ", '<Default Extension="jpg" ContentType="image/jpeg"/><Default ', 1)
        tmp = path + ".tmp"
        with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = ct.encode() if item.filename == "[Content_Types].xml" else zin.read(item.filename)
                zout.writestr(item, data)
        zin.close()
        shutil.move(tmp, path)


def strip_spacing(root):
    """Remove letter-spacing anywhere (e.g. layout kicker placeholder): it breaks Persian letter joining."""
    for el in root.iter():
        if "spc" in el.attrib:
            del el.attrib["spc"]


def main(src, dst):
    prs = Presentation(src)
    for layout in prs.slide_layouts:
        do_shapes(layout.shapes)
        strip_spacing(layout._element)
    for slide in prs.slides:
        do_shapes(slide.shapes)
        strip_spacing(slide._element)
    prs.save(dst)
    fix_content_types(dst)
    # report runs that are still English prose (for QA)
    left = set()
    for slide in Presentation(dst).slides:
        for sh in slide.shapes:
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        if re.search(r"[a-z]{4,} [a-z]{3,}", r.text) and not PERSIAN.search(r.text):
                            left.add(r.text[:70])
    print(dst, "| English-only runs left:", len(left))
    for t in sorted(left)[:40]:
        print("   ", t)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
