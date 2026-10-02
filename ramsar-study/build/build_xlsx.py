# -*- coding: utf-8 -*-
"""Build Ramsar_Database.xlsx from data.py."""
import os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import data as D

OUT = os.path.join(os.path.dirname(__file__), "..", "Ramsar_Database.xlsx")
HEAD = PatternFill("solid", fgColor="0F4C5C")
THIN = Side(style="thin", color="B7C4C2")
FONT = "Vazirmatn"


def qlabel(code):
    e, name, _ = D.Q[code]
    return f"{e} {name}"


def scope_label(s):
    return f"{s} — {D.SCOPE.get(s.split(' ')[0], '')}" if s in D.SCOPE else s


def sheet(wb, title, headers, rows, widths=None, qcol=None, note=None):
    ws = wb.create_sheet(title)
    ws.sheet_view.rightToLeft = True
    r0 = 1
    if note:
        ws.cell(row=1, column=1, value=note).font = Font(name=FONT, italic=True, size=9, color="555555")
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        r0 = 2
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=r0, column=c, value=h)
        cell.fill = HEAD
        cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, row in enumerate(rows, r0 + 1):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=i, column=c, value=v)
            cell.font = Font(name=FONT, size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
        if qcol is not None:
            code = rows[i - r0 - 1][qcol]
            for k, (cd, (e, n, color)) in enumerate(D.Q.items()):
                if code == f"{e} {n}":
                    ws.cell(row=i, column=qcol + 1).fill = PatternFill("solid", fgColor=color)
    for c in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(c)].width = (widths or {}).get(c, 18)
    ws.freeze_panes = ws.cell(row=r0 + 1, column=1)
    ws.auto_filter.ref = f"A{r0}:{get_column_letter(len(headers))}{r0 + len(rows)}"
    return ws


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    ws.sheet_view.rightToLeft = True
    lines = [
        "پایگاه داده مطالعات شهری رامسر — دور دوم پژوهش (مهر 1405)",
        "گروه سوم: مشکات عروجی، محمد امین صادقی، فاطمه صفری‌نژاد، آرینه مارگوسیان، حدیثه صیفوری",
        "درس: مبانی برنامه‌ریزی فضاهای شهری",
        "زنجیره: DATA → ANALYSIS → PROBLEM → NEED → STRATEGY → PROGRAM → DESIGN",
        "",
        "سطح اعتبار داده:",
        *[f"   {e} {n}" for e, n, _ in D.Q.values()],
        "",
        "قانون: هر عدد = منبع + سال + محدوده جغرافیایی (City/County/District/Province) + سطح اعتبار.",
        "شهر رامسر (City) و شهرستان رامسر (County) هرگز با هم جمع یا جایگزین نشده‌اند.",
        "کدهای P01/P02 = فایل‌های پروژه؛ N01… = منابع جدید این دور؛ S01… = کدهای داخل گزارش پایه.",
        f"تاریخ دسترسی به منابع وب: {D.ACCESS_DATE}",
        "محدودیت: در محیط اجرای این پژوهش، دسترسی مستقیم به OSM/Overpass، amar.org.ir، ramsar.ir و ویکی‌پدیا مسدود بود؛",
        "منابع از طریق موتور جستجو خوانده شدند. نقشه‌های خدمات با اسکریپت gis/ramsar_osm_services.py روی سیستم دارای اینترنت تولید می‌شوند.",
    ]
    for i, t in enumerate(lines, 1):
        ws.cell(row=i, column=1, value=t).font = Font(name=FONT, size=11, bold=(i == 1))
    ws.column_dimensions["A"].width = 120

    sheet(wb, "SOURCES",
          ["کد", "سازمان / نویسنده", "عنوان", "سال", "URL", "نوع منبع", "رتبه اعتبار (1 بهترین؛ 0 = سند پروژه)", "توضیح"],
          [list(s) for s in D.SOURCES], {1: 6, 2: 30, 3: 50, 4: 14, 5: 45, 6: 18, 7: 12, 8: 40})

    sheet(wb, "POPULATION",
          ["سال", "Geographic Scope", "جمعیت", "خانوار", "بعد خانوار", "رشد سالانه ٪", "منبع", "اعتبار", "توضیح"],
          [[y, scope_label(s), p if p else "DATA GAP", h if h else "DATA GAP", hs if hs else "—",
            g if g is not None else ("پایه" if p else "—"), src, qlabel(q), n] for y, s, p, h, hs, g, src, q, n in D.POP],
          {1: 8, 2: 20, 9: 50}, qcol=7,
          note="رشد = (Pt/P0)^(1/n) − 1 ؛ بعد خانوار = جمعیت ÷ خانوار")

    sheet(wb, "PER_CAPITA",
          ["شاخص", "مقدار موجود", "جمعیت", "سرانه / شاخص", "استاندارد", "کمبود/مازاد", "سال", "منبع", "اعتبار", "تحلیل کوتاه"],
          [[a, b, c, d, e, f, g, h, qlabel(q), z] for a, b, c, d, e, f, g, h, q, z in D.PER_CAPITA],
          {1: 34, 2: 30, 4: 22, 6: 22, 10: 60}, qcol=8,
          note=f"جمعیت مبنا: شهر رامسر 35,997 نفر (سرشماری 1395، 🟢). سناریوی 1405 فقط برای حساسیت (🟠).")

    sheet(wb, "STANDARDS",
          ["شاخص", "مقدار", "واحد", "مرجع", "سال", "محدوده کاربرد", "اعتبار", "توضیح / دلیل انتخاب"],
          [[a, b, c, d, e, f, qlabel(q), n] for a, b, c, d, e, f, q, n in D.STANDARDS],
          {1: 28, 2: 30, 4: 45, 8: 50}, qcol=6)

    sheet(wb, "SERVICES",
          ["گروه", "نام", "نوع", "موقعیت", "مساحت/ظرفیت", "سال", "Geographic Scope", "منبع", "اعتبار"],
          [[a, b, c, d, e, f, g, h, qlabel(q)] for a, b, c, d, e, f, g, h, q in D.SERVICES],
          {2: 40, 4: 30, 5: 30}, qcol=8)

    sheet(wb, "TOURISM",
          ["شاخص", "مقدار", "واحد", "سال", "Geographic Scope", "منبع", "اعتبار", "توضیح"],
          [[a, b, c, d, e, f, qlabel(q), n] for a, b, c, d, e, f, q, n in D.TOURISM],
          {1: 34, 2: 40, 8: 50}, qcol=6)

    sheet(wb, "HISTORY",
          ["سال", "رویداد", "اثر کالبدی", "اثر اجتماعی", "اثر اقتصادی", "مکان", "منبع", "اعتبار"],
          [[a, b, c, d, e, f, g, qlabel(q)] for a, b, c, d, e, f, g, q in D.HISTORY],
          {1: 18, 2: 50, 3: 30, 4: 25, 5: 25}, qcol=7)

    sheet(wb, "GEOGRAPHY",
          ["شاخص", "مقدار", "سال", "Geographic Scope", "منبع", "اعتبار", "توضیح"],
          [[a, b, c, d, e, qlabel(q), n] for a, b, c, d, e, q, n in D.GEOGRAPHY],
          {1: 34, 2: 45, 7: 50}, qcol=5)

    sheet(wb, "ECONOMY",
          ["شاخص", "مقدار", "سال", "Geographic Scope", "منبع", "اعتبار", "توضیح"],
          [[a, b, c, d, e, qlabel(q), n] for a, b, c, d, e, q, n in D.ECONOMY],
          {1: 34, 2: 45, 7: 50}, qcol=5)

    sheet(wb, "COMPARISON",
          ["شهر", "استان", "جمعیت 1395", "Geographic Scope", "سرانه فضای سبز", "ویژگی گردشگری", "منبع", "اعتبار", "قابلیت مقایسه"],
          [[a, b, c, d, e, f, g, qlabel(q), z] for a, b, c, d, e, f, g, q, z in D.COMPARE],
          {5: 30, 9: 34}, qcol=7)

    sheet(wb, "CONFLICTS",
          ["موضوع", "عدد A", "عدد B", "علت اختلاف", "تصمیم", "اعتبار"],
          [[a, b, c, d, e, qlabel(q)] for a, b, c, d, e, q in D.POP_CONFLICTS] + [
              ["مساحت باغ ملی", "30 ha (N03)", "33 ha (N04)", "محدوده باغ یا گردکردن؛ هیچ‌کدام سند شهرداری نیست", "هر دو گزارش شد؛ بازه 8.33–9.17 m²/نفر", qlabel("CON")],
              ["تخت بیمارستان امام سجاد", "200 (سایت رسمی)", "170 / 197 (فهرست‌ها)", "تخت مصوب در برابر فعال", "200 مصوب؛ بازه 2.30–2.70 تخت/1000 (شهرستان)", qlabel("CON")],
              ["تراز دریای خزر", "−27.5 m (2010)", "< −29 m (2025)", "سال متفاوت؛ احتمال تفاوت مبنای ارتفاعی", "عدد 2025 برای وضع موجود؛ مبنا کنترل شود", qlabel("CON")],
              ["پسماند روزانه", "80 t (میانگین) / 120 t (فصل)", "100–140 t", "منابع و سال‌های مختلف (1395–1404)", "80→120 برای شاخص فصلی", qlabel("CON")],
              ["سال افتتاح هتل قدیم", "1311", "1312", "اختلاف منابع ثانویه", "1311/1312 ثبت شد", qlabel("CON")],
          ], {1: 26, 2: 30, 3: 30, 4: 45, 5: 45}, qcol=5)

    sheet(wb, "LYNCH", ["عنصر", "مصادیق در رامسر", "ارزیابی", "منبع"], [list(r) for r in D.LYNCH], {1: 18, 2: 55, 3: 50})

    sheet(wb, "PROBLEMS", ["کد", "حوزه", "داده", "مسئله", "شاهد (منبع)", "پیامد", "فرصت"],
          [list(r) for r in D.PROBLEMS], {3: 45, 4: 40, 6: 35, 7: 35})
    sheet(wb, "NEEDS", ["مسئله", "نیاز شهری", "پاسخ برنامه‌ریزی"], [list(r) for r in D.NEEDS], {2: 40, 3: 80})
    sheet(wb, "STRATEGIES", ["نوع", "راهبرد", "اقدام کلیدی"], [list(r) for r in D.STRATEGIES], {2: 55, 3: 70})
    sheet(wb, "PROGRAM", ["گروه", "جزء برنامه", "اندازه پیشنهادی", "کاربران", "نیاز مرتبط", "مبنای داده"],
          [list(r) for r in D.PROGRAM], {2: 60, 3: 30}, note="محدوده طراحی پیشنهادی: محور میراثی–ساحلی معلم (هتل قدیم ← ساحل)؛ اندازه‌ها پیشنهاد طراحی‌اند نه داده")
    sheet(wb, "MAPS", ["نقشه", "عنوان", "لایه‌ها", "تگ‌های OSM", "وضعیت"], [list(r) for r in D.MAPS], {3: 45, 4: 55, 5: 30})
    sheet(wb, "SHEETS_A2", ["شیت", "عنوان", "محتوا"], [[a, b, " | ".join(c)] for a, b, c in D.SHEETS], {3: 120})
    sheet(wb, "SUMMARY", ["#", "10 مسئله اصلی", "10 ظرفیت اصلی", "10 نیاز اصلی"],
          [[i + 1, a, b, c] for i, (a, b, c) in enumerate(zip(D.TOP10_ISSUES, D.TOP10_CAPACITIES, D.TOP10_NEEDS))], {2: 55, 3: 55, 4: 50})

    gaps = [[r[0], r[1], "DATA GAP"] for r in D.PER_CAPITA if r[8] == "GAP"]
    gaps += [[r[0], r[1], "DATA GAP"] for r in D.SERVICES if r[8] == "GAP"]
    gaps += [[r[0], r[1], "DATA GAP"] for r in D.TOURISM + D.ECONOMY if (r[6] if len(r) == 8 else r[5]) == "GAP"]
    gaps += [["جمعیت 1405 شهر", "سرشماری ثبت‌مبنا در جریان", "DATA GAP"],
             ["جدول سرانه مصوبه 1389", "بازیابی نشد", "DATA GAP"],
             ["طرح جامع/تفصیلی رامسر و سال تصویب", "یافت نشد", "DATA GAP"],
             ["ساختار سنی و جنسی 1395 شهر", "جدول تفصیلی مرکز آمار در دسترس نبود", "DATA GAP"]]
    sheet(wb, "DATA_GAPS", ["موضوع", "وضعیت", "برچسب"], gaps, {1: 45, 2: 60})

    wb.save(OUT)
    print("saved", os.path.abspath(OUT))


if __name__ == "__main__":
    main()
